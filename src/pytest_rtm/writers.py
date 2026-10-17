"""Write the matrix as JSON, Markdown, HTML and Excel."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, PackageLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .matrix import FAILED, NO_EVIDENCE, NOT_COVERED, NOT_RUN, VERIFIED, Matrix

STATUS_FILL = {
    VERIFIED: "C8E6C9",
    FAILED: "FFCDD2",
    NOT_RUN: "FFF3C4",
    NOT_COVERED: "FFE0B2",
    NO_EVIDENCE: "FFE0B2",
}

_env = Environment(loader=PackageLoader("pytest_rtm", "templates"), autoescape=select_autoescape(["html", "j2"]))


def write_all(matrix: Matrix, out_dir: Path, title: str) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    generated = datetime.now().isoformat(timespec="seconds")
    return [
        write_html(matrix, out_dir / "rtm.html", title, generated),
        write_xlsx(matrix, out_dir / "rtm.xlsx", title, generated),
        write_json(matrix, out_dir / "rtm.json", title, generated),
        write_markdown(matrix, out_dir / "rtm.md", title),
    ]


def write_json(matrix: Matrix, path: Path, title: str, generated: str) -> Path:
    data = {
        "title": title,
        "generated": generated,
        "coverage_percent": round(matrix.coverage, 1),
        "counts": matrix.counts(),
        "requirements": [
            {
                "id": row.requirement.id,
                "category": row.requirement.category,
                "text": row.requirement.text,
                "verification": row.requirement.verification,
                "evidence": row.requirement.evidence,
                "status": row.status,
                "tests": [{"id": test_id, "outcome": outcome} for test_id, outcome in row.tests],
            }
            for row in matrix.rows
        ],
        "untraced_tests": matrix.untraced_tests,
    }
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return path


def write_markdown(matrix: Matrix, path: Path, title: str) -> Path:
    counts = matrix.counts()
    lines = [
        f"### {title}",
        "",
        f"**{matrix.coverage:.1f}% verified**: "
        + ", ".join(f"{status}: {count}" for status, count in counts.items() if count),
        "",
        "| Requirement | Method | Tests | Status |",
        "|---|---|---:|---|",
    ]
    for row in matrix.rows:
        tests = len(row.tests) if row.requirement.verification == "Test" else row.requirement.evidence
        lines.append(f"| {row.requirement.id} | {row.requirement.verification} | {tests} | {row.status} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_html(matrix: Matrix, path: Path, title: str, generated: str) -> Path:
    html = _env.get_template("rtm.html.j2").render(
        matrix=matrix, counts=matrix.counts(), title=title, generated=generated,
        status_class={VERIFIED: "ok", FAILED: "bad", NOT_RUN: "warn", NOT_COVERED: "gap", NO_EVIDENCE: "gap"},
    )
    path.write_text(html, encoding="utf-8")
    return path


def write_xlsx(matrix: Matrix, path: Path, title: str, generated: str) -> Path:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F4A6D")
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, columns, widths):
        ws.append(columns)
        for col, width in enumerate(widths, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill = header_font, header_fill
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "RTM"
    header(ws, ["Req ID", "Category", "Requirement", "Method", "Verified by", "Status"], [11, 14, 60, 13, 70, 14])
    for row in matrix.rows:
        req = row.requirement
        verified_by = req.evidence if req.verification != "Test" else "\n".join(
            f"{test_id} [{outcome}]" for test_id, outcome in row.tests)
        ws.append([req.id, req.category, req.text, req.verification, verified_by, row.status])
        for cell in ws[ws.max_row]:
            cell.alignment = wrap
        ws.cell(row=ws.max_row, column=6).fill = PatternFill("solid", fgColor=STATUS_FILL[row.status])
    ws.auto_filter.ref = ws.dimensions

    tests = wb.create_sheet("Tests")
    header(tests, ["Test", "Requirements", "Outcome"], [80, 30, 12])
    by_test: dict[str, tuple[list[str], str]] = {}
    for row in matrix.rows:
        for test_id, outcome in row.tests:
            by_test.setdefault(test_id, ([], outcome))[0].append(row.requirement.id)
    for test_id, (req_ids, outcome) in sorted(by_test.items()):
        tests.append([test_id, ", ".join(req_ids), outcome])
    for test_id in matrix.untraced_tests:
        tests.append([test_id, "(untraced)", ""])

    summary = wb.create_sheet("Summary")
    summary.append([title])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Generated", generated])
    summary.append(["Requirements", len(matrix.rows)])
    summary.append(["Verified %", round(matrix.coverage, 1)])
    for status, count in matrix.counts().items():
        summary.append([status, count])
    summary.column_dimensions["A"].width = 18

    wb.save(path)
    return path
